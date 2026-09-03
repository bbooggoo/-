"""
역할(관리자/설계사/기술팀 담당자) + 자동/수동 질의 2트랙 승인 워크플로우 테스트.

- 관리자만 업로드 가능 (item 6)
- 자동 질의: 동일 교정을 대공정 단위로 묶어서 일괄 생성, 기술팀 승인/미승인 한 번 (items 1, 2)
- 수동 질의: 4단계 설계사/기술팀 승인 흐름 + 값 제안시 검증 룰셋 강제 (items 1, 3)
- 담당자별 SUMMARY(대공정 필터) + GCS(SPECIALITY GAS·폐액) 종수 (items 4, 8)
- 기술팀 미해결만 보기 필터 (item 7)
"""
from io import BytesIO

from openpyxl import Workbook


def _make_excel(rows: list[list]):
    """UTILITY(위치/라인/건설코드) + GAS/AIR(유량/성상명) 2단 헤더 엑셀."""
    wb = Workbook()
    ws = wb.active
    columns = [("UTILITY", "위치"), ("UTILITY", "라인"), ("UTILITY", "건설코드"),
               ("GAS/AIR", "유량"), ("GAS/AIR", "성상명")]
    for i, (c, l) in enumerate(columns, start=1):
        ws.cell(1, i, c)
        ws.cell(2, i, l)
    for r, row in enumerate(rows, start=3):
        for i, v in enumerate(row, start=1):
            ws.cell(r, i, v)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _upload(client, major_process_id, rows, headers=None):
    return client.post(
        "/api/spec-sheets/upload",
        data={"major_process_id": major_process_id},
        files={"file": ("s.xlsx", _make_excel(rows), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=headers,
    )


def _get_piping(client):
    return next(d for d in client.get("/api/disciplines").json() if d["code"] == "PIPING")


def test_upload_requires_admin(client):
    """item 6: 제원표 업로드는 관리자만."""
    mps = client.get("/api/major-processes").json()
    cvd = next(m for m in mps if m["code"] == "CVD")
    designer = client.post(
        "/api/owners", json={"name": "박설계", "role": "DESIGNER", "discipline_ids": []}
    ).json()

    # 관리자가 아닌 사용자(설계사)로는 업로드 불가 -> 403
    res_no_admin = client.post(
        "/api/spec-sheets/upload",
        data={"major_process_id": cvd["id"]},
        files={"file": ("s.xlsx", _make_excel([["평택", "P4", "PD000001", 5, "N2"]]),
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers={"X-User-Id": str(designer["id"])},
    )
    assert res_no_admin.status_code == 403


def test_admin_upload_succeeds(client):
    mps = client.get("/api/major-processes").json()
    cvd = next(m for m in mps if m["code"] == "CVD")
    res = _upload(client, cvd["id"], [["평택", "P4", "PD000001", 5, "N2"]])
    assert res.status_code == 200
    assert len(res.json()["created"]) == 1


def test_auto_query_batches_identical_corrections_and_applies_on_approval(client):
    """items 1, 2: 동일 교정(같은 대공정+규칙+필드+TO-BE)은 여러 시트에 걸쳐도 질의 1건으로
    묶이고, 기술팀 승인 한 번으로 대상 전체에 반영된다."""
    mps = client.get("/api/major-processes").json()
    cvd = next(m for m in mps if m["code"] == "CVD")
    tech = client.post(
        "/api/owners", json={"name": "김기술", "role": "TECH_LEAD", "major_process_ids": [cvd["id"]]}
    ).json()

    client.post(
        "/api/validation-rules",
        json={
            "name": "성상명 자동 교정", "rule_type": "alias_correction",
            "major_process_id": cvd["id"], "field_name_pattern": r"^GAS/AIR_성상명$",
            "params": {"aliases": {"GN2": "N2", "LN2": "N2"}}, "severity": "WARN", "active": True,
        },
    )

    up = _upload(client, cvd["id"], [
        ["평택", "P4", "PD000001", 5, "GN2"],
        ["평택", "P4", "PD000002", 5, "LN2"],
    ])
    sheets = up.json()["created"]
    for s in sheets:
        client.post(f"/api/spec-sheets/{s['id']}/validate")

    threads = client.post("/api/qa-threads/generate-auto").json()
    assert len(threads) == 1  # 서로 다른 시트지만 같은 교정이라 하나로 묶임
    thread = threads[0]
    assert thread["query_type"] == "AUTO"
    assert thread["to_be_value"] == "N2"
    assert len(thread["targets"]) == 2

    # 다시 생성해도 이미 배치된 항목은 중복 배치하지 않는다
    threads_again = client.post("/api/qa-threads/generate-auto").json()
    assert threads_again == []

    # 대공정 권한 없는 기술팀은 승인 불가
    other_tech = client.post(
        "/api/owners", json={"name": "이엣치", "role": "TECH_LEAD", "major_process_ids": []}
    ).json()
    denied = client.post(
        f"/api/qa-threads/{thread['id']}/auto-decision",
        json={"approve": True, "decided_by": "이엣치"},
        headers={"X-User-Id": str(other_tech["id"])},
    )
    assert denied.status_code == 403

    approved = client.post(
        f"/api/qa-threads/{thread['id']}/auto-decision",
        json={"approve": True, "decided_by": "김기술"},
        headers={"X-User-Id": str(tech["id"])},
    ).json()
    assert approved["status"] == "RESOLVED"

    for s in sheets:
        detail = client.get(f"/api/spec-sheets/{s['id']}").json()
        gas_field = next(f for f in detail["fields"] if f["field_name"] == "GAS/AIR_성상명")
        assert gas_field["value"] == "N2"
        assert detail["open_issue_count"] == 0
        assert detail["status"] == "RESOLVED"

    corrections = client.get("/api/corrections").json()
    assert len(corrections) == 2
    assert all(c["source"] == "AUTO_SUGGESTED" for c in corrections)


def test_auto_query_rejection_reopens_issue_without_changing_value(client):
    mps = client.get("/api/major-processes").json()
    cvd = next(m for m in mps if m["code"] == "CVD")
    tech = client.post(
        "/api/owners", json={"name": "김기술", "role": "TECH_LEAD", "major_process_ids": [cvd["id"]]}
    ).json()
    client.post(
        "/api/validation-rules",
        json={
            "name": "성상명 자동 교정", "rule_type": "alias_correction",
            "major_process_id": cvd["id"], "field_name_pattern": r"^GAS/AIR_성상명$",
            "params": {"aliases": {"GN2": "N2"}}, "severity": "WARN", "active": True,
        },
    )
    up = _upload(client, cvd["id"], [["평택", "P4", "PD000001", 5, "GN2"]])
    sheet = up.json()["created"][0]
    client.post(f"/api/spec-sheets/{sheet['id']}/validate")
    thread = client.post("/api/qa-threads/generate-auto").json()[0]

    rejected = client.post(
        f"/api/qa-threads/{thread['id']}/auto-decision",
        json={"approve": False, "decided_by": "김기술"},
        headers={"X-User-Id": str(tech["id"])},
    ).json()
    assert rejected["status"] == "REJECTED"

    detail = client.get(f"/api/spec-sheets/{sheet['id']}").json()
    gas_field = next(f for f in detail["fields"] if f["field_name"] == "GAS/AIR_성상명")
    assert gas_field["value"] == "GN2"  # 값은 그대로
    assert detail["open_issue_count"] == 1  # 이슈는 다시 미해결로


def test_manual_query_four_stage_approval_with_validation_ruleset(client):
    """items 1, 3: 수동 질의 OPEN -> TECH_ANSWERED -> ANSWER_APPROVED -> VALUE_PROPOSED
    -> RESOLVED 전 단계와, 값 제안이 검증 룰셋(성상명 허용값)을 통과해야 함을 확인한다."""
    mps = client.get("/api/major-processes").json()
    cvd = next(m for m in mps if m["code"] == "CVD")
    piping = _get_piping(client)
    tech = client.post(
        "/api/owners", json={"name": "김기술", "role": "TECH_LEAD", "major_process_ids": [cvd["id"]]}
    ).json()
    designer = client.post(
        "/api/owners", json={"name": "박설계", "role": "DESIGNER", "discipline_ids": [piping["id"]]}
    ).json()
    client.post(
        "/api/validation-rules",
        json={
            "name": "성상명 표준화", "rule_type": "standardized_enum",
            "major_process_id": cvd["id"], "field_name_pattern": r"^GAS/AIR_성상명$",
            "params": {"allowed_values": ["N2", "O2"]}, "severity": "ERROR", "active": True,
        },
    )

    up = _upload(client, cvd["id"], [["평택", "P4", "PD000001", 5, "N2"]])
    sheet = up.json()["created"][0]
    field = next(f for f in sheet["fields"] if f["field_name"] == "GAS/AIR_성상명")

    thread = client.post(
        f"/api/spec-sheets/{sheet['id']}/qa-threads",
        json={
            "title": "성상명 확인", "discipline_id": piping["id"], "spec_field_id": field["id"],
            "question": "이 값이 맞는지 확인 부탁드립니다.", "author_name": "박설계",
        },
        headers={"X-User-Id": str(designer["id"])},
    ).json()
    assert thread["status"] == "OPEN"

    # OPEN 단계에서는 아직 값 제안 불가
    early = client.post(
        f"/api/qa-threads/{thread['id']}/propose-value",
        json={"proposed_by": "김기술", "new_value": "O2"},
        headers={"X-User-Id": str(tech["id"])},
    )
    assert early.status_code == 400

    answered = client.post(
        f"/api/qa-threads/{thread['id']}/messages",
        json={"author_name": "김기술", "role": "ANSWER", "content": "확인했습니다."},
        headers={"X-User-Id": str(tech["id"])},
    ).json()
    assert answered["status"] == "TECH_ANSWERED"

    # 설계사가 아닌 사람은 답변 승인 불가
    tech_tries_approve = client.post(
        f"/api/qa-threads/{thread['id']}/answer-decision",
        json={"approve": True, "decided_by": "김기술"},
        headers={"X-User-Id": str(tech["id"])},
    )
    assert tech_tries_approve.status_code == 403

    approved = client.post(
        f"/api/qa-threads/{thread['id']}/answer-decision",
        json={"approve": True, "decided_by": "박설계"},
        headers={"X-User-Id": str(designer["id"])},
    ).json()
    assert approved["status"] == "ANSWER_APPROVED"

    # 허용되지 않은 값은 검증 룰셋에 걸려 거부됨
    bad = client.post(
        f"/api/qa-threads/{thread['id']}/propose-value",
        json={"proposed_by": "김기술", "new_value": "XX"},
        headers={"X-User-Id": str(tech["id"])},
    )
    assert bad.status_code == 400
    assert "errors" in bad.json()["detail"]

    proposed = client.post(
        f"/api/qa-threads/{thread['id']}/propose-value",
        json={"proposed_by": "김기술", "new_value": "O2"},
        headers={"X-User-Id": str(tech["id"])},
    ).json()
    assert proposed["status"] == "VALUE_PROPOSED"
    assert proposed["to_be_value"] == "O2"

    # DB는 아직 반영 전
    still_old = client.get(f"/api/spec-sheets/{sheet['id']}").json()
    assert next(f for f in still_old["fields"] if f["field_name"] == "GAS/AIR_성상명")["value"] == "N2"

    final = client.post(
        f"/api/qa-threads/{thread['id']}/final-decision",
        json={"approve": True, "decided_by": "박설계"},
        headers={"X-User-Id": str(designer["id"])},
    ).json()
    assert final["status"] == "RESOLVED"

    updated = client.get(f"/api/spec-sheets/{sheet['id']}").json()
    assert next(f for f in updated["fields"] if f["field_name"] == "GAS/AIR_성상명")["value"] == "O2"

    corrections = client.get("/api/corrections").json()
    assert len(corrections) == 1
    assert corrections[0]["source"] == "MANUAL"
    assert corrections[0]["old_value"] == "N2"
    assert corrections[0]["new_value"] == "O2"


def test_open_only_filter_hides_resolved_threads(client):
    """item 7: 기술팀 화면에서 미해결만 보기."""
    mps = client.get("/api/major-processes").json()
    cvd = next(m for m in mps if m["code"] == "CVD")
    piping = _get_piping(client)
    designer = client.post(
        "/api/owners", json={"name": "박설계", "role": "DESIGNER", "discipline_ids": [piping["id"]]}
    ).json()

    up = _upload(client, cvd["id"], [["평택", "P4", "PD000001", 5, "N2"], ["평택", "P4", "PD000002", 5, "N2"]])
    sheets = up.json()["created"]

    for i, sheet in enumerate(sheets):
        field = next(f for f in sheet["fields"] if f["field_name"] == "GAS/AIR_성상명")
        client.post(
            f"/api/spec-sheets/{sheet['id']}/qa-threads",
            json={
                "title": f"질의 {i}", "discipline_id": piping["id"], "spec_field_id": field["id"],
                "question": "확인 요청", "author_name": "박설계",
            },
            headers={"X-User-Id": str(designer["id"])},
        )

    all_threads = client.get(f"/api/qa-threads?major_process_id={cvd['id']}").json()
    assert len(all_threads) == 2

    open_only = client.get(f"/api/qa-threads?major_process_id={cvd['id']}&open_only=true").json()
    assert len(open_only) == 2  # 아직 둘 다 미해결


def test_summary_filters_by_major_process_and_counts_gcs_materials(client):
    """items 4, 8: 대공정별 SUMMARY + GCS(SPECIALITY GAS+폐액) 종수(중복 자재명 dedup)."""
    mps = client.get("/api/major-processes").json()
    cvd = next(m for m in mps if m["code"] == "CVD")
    etch = next(m for m in mps if m["code"] == "ETCH")

    wb = Workbook()
    ws = wb.active
    columns = [("UTILITY", "위치"), ("UTILITY", "라인"), ("UTILITY", "건설코드"),
               ("SPECIALITY GAS", "자재명"), ("SPECIALITY GAS", "배관수량"),
               ("폐액", "자재명"), ("폐액", "실사용량")]
    for i, (c, l) in enumerate(columns, start=1):
        ws.cell(1, i, c)
        ws.cell(2, i, l)
    rows = [
        ["평택", "P4", "PD000001", "SiH4", 1, None, None],
        ["평택", "P4", "PD000001", "PH3", 1, None, None],
        ["평택", "P4", "PD000001", None, None, "ACID WASTE", 2.0],
        ["평택", "P4", "PD000001", None, None, "SiH4", 1.0],  # 같은 자재명이 양쪽에 있어도 종수는 1건
    ]
    for r, row in enumerate(rows, start=3):
        for i, v in enumerate(row, start=1):
            ws.cell(r, i, v)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    res = client.post(
        "/api/spec-sheets/upload", data={"major_process_id": cvd["id"]},
        files={"file": ("s.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert res.status_code == 200

    _upload(client, etch["id"], [["평택", "P4", "PD000002", 5, "N2"]])  # 다른 대공정 - summary에서 제외돼야 함

    summary = client.get(f"/api/summary?major_process_id={cvd['id']}").json()
    assert summary["sheet_count"] == 1
    assert summary["gcs_material_count"] == 3  # SiH4, PH3, ACID WASTE (중복 제거)
    assert sorted(summary["gcs_materials"]) == ["ACID WASTE", "PH3", "SiH4"]

    cat_by_name = {c["category"]: c for c in summary["categories"]}
    assert cat_by_name["SPECIALITY GAS"]["distinct_count"] == 2
    assert cat_by_name["폐액"]["distinct_count"] == 2

    summary_all = client.get("/api/summary").json()
    assert summary_all["sheet_count"] == 2  # 필터 없으면 전체
