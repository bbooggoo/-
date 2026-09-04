"""
제원표를 원본 CAD 레이아웃(좌표) 그대로 엑셀로 재출력하되,
검증 결과가 있는 셀은 색상으로 표시한다 (ERROR=빨강, WARN=노랑).
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from ..models import RuleSeverity, SpecSheet, ValidationResultStatus

ERROR_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
WARN_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")


def export_spec_sheet(spec_sheet: SpecSheet) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = spec_sheet.sheet_name[:31] if spec_sheet.sheet_name else "Sheet1"

    # 필드별 미해결 검증 결과 중 가장 심각한 것 매핑
    flagged: dict[int, RuleSeverity] = {}
    for vr in spec_sheet.validation_results:
        if vr.status == ValidationResultStatus.RESOLVED or vr.spec_field_id is None:
            continue
        current = flagged.get(vr.spec_field_id)
        if current is None or vr.severity == RuleSeverity.ERROR:
            flagged[vr.spec_field_id] = vr.severity

    for f in spec_sheet.fields:
        cell = ws.cell(row=f.row_index, column=f.col_index, value=f.value)
        severity = flagged.get(f.id)
        if severity == RuleSeverity.ERROR:
            cell.fill = ERROR_FILL
        elif severity == RuleSeverity.WARN:
            cell.fill = WARN_FILL

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
